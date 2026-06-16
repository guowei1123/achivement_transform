from pathlib import Path

import pypdfium2 as pdfium
from PIL import Image
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "output"
OUT.mkdir(exist_ok=True)

TEAL = "69AFAC"
LIGHT_GRAY = "F2F2F2"
TEXT_GRAY = "666666"
BLACK = "000000"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge in kwargs:
            edge_data = kwargs.get(edge)
            tag = "w:{}".format(edge)
            element = borders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                borders.append(element)
            for key in ["sz", "val", "color", "space"]:
                if key in edge_data:
                    element.set(qn("w:{}".format(key)), str(edge_data[key]))


def set_table_borders(table, color="FFFFFF", size="0"):
    for row in table.rows:
        for cell in row.cells:
            set_cell_border(
                cell,
                top={"val": "single", "sz": size, "color": color},
                bottom={"val": "single", "sz": size, "color": color},
                left={"val": "single", "sz": size, "color": color},
                right={"val": "single", "sz": size, "color": color},
            )


def set_width(cell, width_cm):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_cm * 567)))
    tc_w.set(qn("w:type"), "dxa")


def set_run_font(run, size=None, bold=False, color=BLACK, name="Arial"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    if size:
        run.font.size = Pt(size)


def para(cell_or_doc, text="", size=10, bold=False, color=BLACK, align=None, before=0, after=2, name="Arial"):
    p = cell_or_doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.05
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    set_run_font(r, size=size, bold=bold, color=color, name=name)
    return p


def clear_cell(cell):
    for p in cell.paragraphs:
        p._element.getparent().remove(p._element)


def extract_photo():
    photo_path = OUT / "guo_wei_photo.png"
    if photo_path.exists():
        return photo_path
    pdf = pdfium.PdfDocument(str(ROOT / "source_resume.pdf"))
    page = pdf[0]
    image = page.render(scale=2.0).to_pil()
    # Crop from the rendered original resume. Coordinates are for the 2x A4 render.
    crop = image.crop((955, 55, 1132, 284))
    crop.save(photo_path)
    return photo_path


def setup_doc(margins_cm=(1.45, 1.35, 1.2, 1.35)):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    top, right, bottom, left = margins_cm
    section.top_margin = Cm(top)
    section.right_margin = Cm(right)
    section.bottom_margin = Cm(bottom)
    section.left_margin = Cm(left)
    for style_name in ["Normal", "Body Text"]:
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(10)
    return doc


def section_heading(doc, title):
    tbl = doc.add_table(rows=1, cols=2)
    tbl.autofit = False
    set_table_borders(tbl)
    left, right = tbl.rows[0].cells
    set_width(left, 4.5)
    set_width(right, 13.3)
    set_cell_shading(left, TEAL)
    left.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    clear_cell(left)
    p = left.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_run_font(r, size=13, bold=False, color="FFFFFF")
    set_cell_border(right, bottom={"val": "single", "sz": "18", "color": TEAL})
    clear_cell(right)
    right.add_paragraph("")


def role_block(doc, date, org, role, lines, date_size=10):
    top = doc.add_table(rows=1, cols=2)
    top.autofit = False
    set_table_borders(top)
    left, right = top.rows[0].cells
    set_width(left, 7.0)
    set_width(right, 10.8)
    clear_cell(left)
    clear_cell(right)
    para(left, date, size=date_size, bold=True, after=0)
    p = right.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(org)
    set_run_font(r, size=10, bold=True)
    if role:
        para(doc, role, size=10, color=TEXT_GRAY, after=1)
    for line in lines:
        para(doc, line, size=8.5, color=TEXT_GRAY, after=1)


def build_resume():
    doc = setup_doc()
    photo = extract_photo()

    header = doc.add_table(rows=1, cols=2)
    header.autofit = False
    set_table_borders(header)
    info_cell, photo_cell = header.rows[0].cells
    set_width(info_cell, 15.0)
    set_width(photo_cell, 2.8)
    clear_cell(info_cell)
    clear_cell(photo_cell)
    p = info_cell.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("GUO Wei")
    set_run_font(r, size=24, bold=False)

    info = info_cell.add_table(rows=2, cols=3)
    info.autofit = False
    set_table_borders(info, color=LIGHT_GRAY, size="4")
    for row in info.rows:
        for cell in row.cells:
            set_cell_shading(cell, LIGHT_GRAY)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    items = [
        ("Gender:", "Male"),
        ("Age:", "20"),
        ("Phone:", "15074714635"),
        ("Hometown:", "Hengyang, Hunan"),
        ("Email:", "326954@whut.edu.cn"),
        ("Political:", "Prob. CPC Member"),
    ]
    for cell, (k, v) in zip([c for row in info.rows for c in row.cells], items):
        clear_cell(cell)
        p = cell.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        rk = p.add_run(k + " ")
        set_run_font(rk, size=8.8, bold=True, color="444444")
        rv = p.add_run(v)
        set_run_font(rv, size=8.8, color="444444")
    p = photo_cell.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.add_run().add_picture(str(photo), width=Cm(2.35))

    section_heading(doc, "Education")
    role_block(
        doc,
        "2021.9-2025.6",
        "Wuhan University of Technology",
        "E-Commerce | Bachelor",
        [
            "Honors: Model Merit Student, Outstanding Undergraduate Graduate, National Encouragement Scholarship; compulsory-course GPA 4.153/5; ranked 3/51 in major.",
        ],
    )
    role_block(
        doc,
        "2025.9-2028.6",
        "Wuhan University of Technology",
        "International Business | Master",
        ["Awarded the First-Class Graduate Entrance Scholarship."],
    )

    section_heading(doc, "Project Experience")
    role_block(
        doc,
        "2025.12-2026.4",
        "Intelligent Technology Service Research Integrating Knowledge Graphs and AI LLMs",
        "",
        [
            "Built an LLM- and knowledge-graph-based intelligent service system for technology transfer in key sectors such as shipping, enabling precise matching between research outputs and industry demand.",
            "Collected and cleaned data across 20+ upstream and downstream shipping sub-sectors, building the underlying dataset for an industrial knowledge graph with 50,000+ entities and relationships.",
            "Participated in LangChain-based Agent development; used Codex to support selected implementation tasks and optimized LLM calls, prompt design, and matching logic, reaching 85% Top-5 matching accuracy in test samples.",
        ],
    )
    role_block(
        doc,
        "2026.1-2026.6",
        "Wuling Motors Patent Service Agent Development",
        "",
        [
            "Developed a patent-service intelligent agent for Guangxi Liuzhou Wuling Motors to support patent disclosure drafting, search-query generation, and patent search report generation.",
            "Mapped core needs of IPR and PE engineers across patent submission, disclosure drafting, novelty search, and report generation; helped define functional modules and business workflows.",
            "Designed and debugged five prompt templates and generation rules for mechanical structure, computer process, and industrial design scenarios, reducing initial disclosure drafting time from about two days to around ten minutes.",
        ],
    )

    section_heading(doc, "Campus Experience")
    role_block(
        doc,
        "2025.10-2026.7",
        "MPA Student Office Assistant",
        "",
        [
            "Managed teaching materials for 200+ MPA students and archived 500+ examination papers.",
            "Supported coordination of MPA admission interview volunteer activities, helping ensure smooth interview participation for 200+ candidates.",
        ],
    )

    section_heading(doc, "Relevant Skills")
    skills = [
        "Programming and databases: Python, Java, MySQL, PostgreSQL.",
        "AI tools and concepts: LLM, RAG, Agent, Prompt Engineering; hands-on experience with Coze, LangChain, and Codex.",
        "Research and productivity tools: Word, Excel, PowerPoint, Tableau; able to structure data, documents, and presentations for research delivery.",
        "Language and certificates: CET-6; C2 driving license.",
    ]
    for i, skill in enumerate(skills, 1):
        para(doc, f"{i}. {skill}", size=8.6, color=TEXT_GRAY, after=0)

    out = OUT / "Guo_Wei_English_Resume.docx"
    doc.save(out)
    return out


def build_cover_letter():
    doc = setup_doc(margins_cm=(1.8, 2.0, 1.8, 2.0))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("郭伟")
    set_run_font(r, size=18, bold=True, color=TEAL, name="Microsoft YaHei")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(16)
    r = p.add_run("326954@whut.edu.cn | 15074714635 | 武汉理工大学")
    set_run_font(r, size=9, color=TEXT_GRAY, name="Microsoft YaHei")

    for line in [
        "2026年6月16日",
        "",
        "Z Partners 招聘团队：",
        "",
    ]:
        para(doc, line, size=10.5, color=BLACK, after=2, name="Microsoft YaHei")

    paragraphs = [
        "您好！我申请 Z Partners 长期实习生岗位。我目前就读于武汉理工大学国际商务硕士项目，本科专业为电子商务，曾获得国家励志奖学金、校三好学生标兵、校优秀本科毕业生等荣誉，专业排名 3/51。Z Partners 强调扎实研究、真实洞察和面向投资决策的定制化服务，这与我一直训练自己的方向高度一致：在复杂信息中梳理结构，在证据基础上形成判断。",
        "我近期的项目经历集中在产业研究和 AI 应用落地的交叉领域。在“融合知识图谱和 AI 大模型的智能科技服务研究”中，我围绕航运等重点产业，完成 20+ 个上下游细分领域的数据抓取与清洗，参与构建包含 5 万+ 实体/关系的产业知识图谱底层数据集；同时基于 LangChain 参与智能体 Agent 搭建，优化大模型调用、Prompt 设计与匹配逻辑，在测试样本中实现 Top-5 匹配准确率 85%。这段经历让我意识到，研究工作不能停留在资料堆积，而要先拆清产业链上下游、关键参与者和供需关系，再通过多来源信息交叉验证，提炼出真正能解释业务问题的结构化结论。",
        "在项目推进过程中，我也进一步理解了 AI 商业化落地的判断逻辑：技术方案是否有价值，关键不只在模型能力本身，而在于它能否嵌入具体业务流程、解决明确痛点并带来可衡量的效率提升。因此，我在整理行业数据和参与 Agent 搭建时，会主动思考不同环节的真实需求、匹配逻辑背后的商业假设，以及技术工具对产业服务模式可能产生的影响。这种从信息收集到结构化分析、再到形成判断的训练，与咨询和投资研究中“把复杂信息转化为有效洞察”的要求高度相关。",
        "我对 Z Partners 的兴趣，来自于岗位所要求的独立研究能力和主动思考意识。我熟悉 Excel、PPT、Word、Python、SQL 等工具，也具备 Coze、Codex 等 AI 工具的实践经验。更重要的是，我愿意持续追问市场为什么如此运行、产业价值在哪里产生、哪些证据足以支撑判断，并希望在真实咨询与投资研究项目中继续打磨这种研究能力。",
        "如果有机会加入 Z Partners，我希望能在相关项目中贡献自己的研究执行、信息结构化和技术理解能力。我近期可以到岗，并能够按照岗位要求持续实习至 2026 年暑期结束，每周出勤至少 3 天。",
        "感谢您审阅我的申请，期待有机会进一步交流。",
        "",
        "此致",
        "敬礼！",
        "郭伟",
    ]
    for text in paragraphs:
        para(doc, text, size=10.5, color=BLACK, after=8, name="Microsoft YaHei")

    out = OUT / "Guo_Wei_ZPartners_Chinese_Cover_Letter.docx"
    doc.save(out)
    return out


if __name__ == "__main__":
    print(build_resume())
    print(build_cover_letter())


def draw_wrapped(c, text, x, y, width, style):
    p = Paragraph(text, style)
    _, h = p.wrap(width, 1000)
    p.drawOn(c, x, y - h)
    return y - h


def build_resume_pdf():
    photo = extract_photo()
    path = OUT / "Guo_Wei_English_Resume.pdf"
    c = canvas.Canvas(str(path), pagesize=A4)
    w, h = A4
    teal = colors.HexColor("#" + TEAL)
    gray = colors.HexColor("#" + TEXT_GRAY)
    light = colors.HexColor("#" + LIGHT_GRAY)

    body = ParagraphStyle("body", fontName="Helvetica", fontSize=8.2, leading=10.5, textColor=gray, alignment=TA_LEFT)
    small = ParagraphStyle("small", fontName="Helvetica", fontSize=7.25, leading=9.0, textColor=gray, alignment=TA_LEFT)
    bold = ParagraphStyle("bold", fontName="Helvetica-Bold", fontSize=9.5, leading=11, textColor=colors.black)
    right_bold = ParagraphStyle("right_bold", parent=bold, alignment=TA_RIGHT)

    c.setFillColor(teal)
    c.rect(0, 0, 14 * mm, h, fill=1, stroke=0)
    x0 = 48
    y = h - 42
    c.setFillColor(colors.black)
    c.setFont("Helvetica", 24)
    c.drawString(x0, y, "GUO Wei")

    info_y = y - 52
    c.setFillColor(light)
    c.rect(x0, info_y, 432, 50, fill=1, stroke=0)
    items = [
        ("Gender:", "Male"), ("Age:", "20"), ("Phone:", "15074714635"),
        ("Hometown:", "Hengyang, Hunan"), ("Email:", "326954@whut.edu.cn"), ("Political:", "Prob. CPC Member"),
    ]
    col_w = 144
    for idx, (k, v) in enumerate(items):
        row = idx // 3
        col = idx % 3
        tx = x0 + 12 + col * col_w
        ty = info_y + 33 - row * 24
        c.setFont("Helvetica-Bold", 8.5)
        c.setFillColor(colors.HexColor("#444444"))
        c.drawString(tx, ty, k)
        c.setFont("Helvetica", 8.5)
        c.drawString(tx + c.stringWidth(k, "Helvetica-Bold", 8.5) + 4, ty, v)
    c.drawImage(str(photo), w - 108, h - 142, width=78, height=100, preserveAspectRatio=True, mask="auto")

    def heading(title, y_pos):
        c.setFillColor(teal)
        c.rect(x0, y_pos - 18, 120, 22, fill=1, stroke=0)
        c.setFillColor(colors.white)
        c.setFont("Helvetica", 13)
        c.drawCentredString(x0 + 60, y_pos - 12, title)
        c.setStrokeColor(teal)
        c.setLineWidth(2)
        c.line(x0 + 120, y_pos - 18, w - 35, y_pos - 18)
        return y_pos - 38

    def block(date, org, role, lines, y_pos):
        c.setFillColor(colors.black)
        c.setFont("Helvetica-Bold", 10)
        c.drawString(x0, y_pos, date)
        p = Paragraph(org, right_bold)
        _, ph = p.wrap(260, 50)
        p.drawOn(c, w - 310, y_pos - ph + 2)
        y2 = y_pos - max(16, ph + 8)
        if role:
            y2 = draw_wrapped(c, role, x0, y2 + 3, w - x0 - 35, body) - 1
        for line in lines:
            y2 = draw_wrapped(c, line, x0, y2, w - x0 - 35, small) - 2
        return y2 - 22

    y = info_y - 36
    y = heading("Education", y)
    y = block("2021.9-2025.6", "Wuhan University of Technology", "E-Commerce | Bachelor", ["Honors: Model Merit Student, Outstanding Undergraduate Graduate, National Encouragement Scholarship; compulsory-course GPA 4.153/5; ranked 3/51 in major."], y)
    y = block("2025.9-2028.6", "Wuhan University of Technology", "International Business | Master", ["Awarded the First-Class Graduate Entrance Scholarship."], y)

    y = heading("Project Experience", y - 8)
    y = block("2025.12-2026.4", "Intelligent Technology Service Research Integrating Knowledge Graphs and AI LLMs", "", [
        "Project: Built an LLM- and knowledge-graph-based intelligent service system for technology transfer in key sectors such as shipping, enabling precise matching between research outputs and industry demand.",
        "Responsibilities: Collected and cleaned data across 20+ upstream and downstream shipping sub-sectors, building the underlying dataset for an industrial knowledge graph with 50,000+ entities and relationships.",
        "Participated in LangChain-based Agent development; used Codex to support selected implementation tasks and optimized LLM calls, prompt design, and matching logic, reaching 85% Top-5 matching accuracy in test samples.",
    ], y)
    y = block("2026.1-2026.6", "Wuling Motors Patent Service Agent Development", "", [
        "Project: Developed a patent-service intelligent agent for Guangxi Liuzhou Wuling Motors to support patent disclosure drafting, search-query generation, and patent search report generation.",
        "Responsibilities: Mapped core needs of IPR and PE engineers across patent submission, disclosure drafting, novelty search, and report generation; helped define functional modules and business workflows.",
        "Designed and debugged five prompt templates and generation rules, reducing initial disclosure drafting time from about two days to around ten minutes.",
    ], y)

    y = heading("Campus Experience", y - 8)
    y = block("2025.10-2026.7", "MPA Student Office Assistant", "", [
        "Managed teaching materials for 200+ MPA students and archived 500+ examination papers.",
        "Supported coordination of MPA admission interview volunteer activities, helping ensure smooth interview participation for 200+ candidates.",
    ], y)

    y = heading("Relevant Skills", y - 8)
    for i, skill in enumerate([
        "Programming and databases: Python, Java, MySQL, PostgreSQL.",
        "AI tools and concepts: LLM, RAG, Agent, Prompt Engineering; hands-on experience with Coze, LangChain, and Codex.",
        "Research and productivity tools: Word, Excel, PowerPoint, Tableau; able to structure data, documents, and presentations for research delivery.",
        "Language and certificates: CET-6; C2 driving license.",
    ], 1):
        y = draw_wrapped(c, f"{i}. {skill}", x0, y + 3, w - x0 - 35, small) - 1

    c.save()
    return path


def build_cover_letter_pdf():
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    path = OUT / "Guo_Wei_ZPartners_Chinese_Cover_Letter.pdf"
    c = canvas.Canvas(str(path), pagesize=A4)
    w, h = A4
    teal = colors.HexColor("#" + TEAL)
    body = ParagraphStyle("letter", fontName="STSong-Light", fontSize=10.8, leading=17, textColor=colors.black, alignment=TA_LEFT)
    x = 62
    y = h - 55
    c.setFillColor(teal)
    c.setFont("STSong-Light", 20)
    c.drawCentredString(w / 2, y, "郭伟")
    y -= 18
    c.setFillColor(colors.HexColor("#666666"))
    c.setFont("STSong-Light", 9)
    c.drawCentredString(w / 2, y, "326954@whut.edu.cn | 15074714635 | 武汉理工大学")
    y -= 32
    text_parts = [
        "2026年6月16日<br/><br/>Z Partners 招聘团队：",
        "您好！我申请 Z Partners 长期实习生岗位。我目前就读于武汉理工大学国际商务硕士项目，本科专业为电子商务，曾获得国家励志奖学金、校三好学生标兵、校优秀本科毕业生等荣誉，专业排名 3/51。Z Partners 强调扎实研究、真实洞察和面向投资决策的定制化服务，这与我一直训练自己的方向高度一致：在复杂信息中梳理结构，在证据基础上形成判断。",
        "我近期的项目经历集中在产业研究和 AI 应用落地的交叉领域。在“融合知识图谱和 AI 大模型的智能科技服务研究”中，我围绕航运等重点产业，完成 20+ 个上下游细分领域的数据抓取与清洗，参与构建包含 5 万+ 实体/关系的产业知识图谱底层数据集；同时基于 LangChain 参与智能体 Agent 搭建，优化大模型调用、Prompt 设计与匹配逻辑，在测试样本中实现 Top-5 匹配准确率 85%。这段经历让我意识到，研究工作不能停留在资料堆积，而要先拆清产业链上下游、关键参与者和供需关系，再通过多来源信息交叉验证，提炼出真正能解释业务问题的结构化结论。",
        "在项目推进过程中，我也进一步理解了 AI 商业化落地的判断逻辑：技术方案是否有价值，关键不只在模型能力本身，而在于它能否嵌入具体业务流程、解决明确痛点并带来可衡量的效率提升。因此，我在整理行业数据和参与 Agent 搭建时，会主动思考不同环节的真实需求、匹配逻辑背后的商业假设，以及技术工具对产业服务模式可能产生的影响。这种从信息收集到结构化分析、再到形成判断的训练，与咨询和投资研究中“把复杂信息转化为有效洞察”的要求高度相关。",
        "我对 Z Partners 的兴趣，来自于岗位所要求的独立研究能力和主动思考意识。我熟悉 Excel、PPT、Word、Python、SQL 等工具，也具备 Coze、Codex 等 AI 工具的实践经验。更重要的是，我愿意持续追问市场为什么如此运行、产业价值在哪里产生、哪些证据足以支撑判断，并希望在真实咨询与投资研究项目中继续打磨这种研究能力。",
        "如果有机会加入 Z Partners，我希望能在相关项目中贡献自己的研究执行、信息结构化和技术理解能力。我近期可以到岗，并能够按照岗位要求持续实习至 2026 年暑期结束，每周出勤至少 3 天。",
        "感谢您审阅我的申请，期待有机会进一步交流。<br/><br/>此致<br/>敬礼！<br/><br/>郭伟",
    ]
    for text in text_parts:
        y = draw_wrapped(c, text, x, y, w - 2 * x, body) - 10
    c.save()
    return path


if __name__ == "__main__":
    print(build_resume_pdf())
    print(build_cover_letter_pdf())
